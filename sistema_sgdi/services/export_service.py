"""
Serviço de Exportação Corporativa Multiformato do SGDI.
Gera relatórios em CSV (utf-8-sig com delimitador ;), Excel estilizado (.xlsx com openpyxl)
e PDF corporativo (com markdown_pdf), todos contendo cabeçalho institucional, rodapé com filtros e timestamp.
"""

import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from markdown_pdf import MarkdownPdf, Section
from flask import Response, send_file
from config import NOME_EMPRESA, NOME_ORGANIZACAO


def gerar_csv_response(metricas):
    """Gera resposta HTTP com arquivo CSV formatado com cabeçalho, dados e rodapé."""
    demandas = metricas['demandas_lista']
    sio = io.StringIO()
    writer = csv.writer(sio, delimiter=';', quoting=csv.QUOTE_MINIMAL)

    # Cabeçalho institucional
    writer.writerow([f'{NOME_EMPRESA} - {NOME_ORGANIZACAO}'])
    writer.writerow(['RELATÓRIO ANALÍTICO DE DEMANDAS E SLA'])
    writer.writerow([])

    # Resumo de KPIs
    writer.writerow(['--- RESUMO DE MÉTRICAS GERAIS ---'])
    writer.writerow(['Total de Demandas', metricas['total_demandas']])
    writer.writerow(['Abertas', metricas['total_abertas']])
    writer.writerow(['Em Andamento', metricas['total_em_andamento']])
    writer.writerow(['Concluídas', metricas['total_concluidas']])
    writer.writerow(['Canceladas', metricas['total_canceladas']])
    writer.writerow(['Fora do SLA / Atrasadas', metricas['total_fora_sla']])
    writer.writerow(['Tempo Médio de Resolução', metricas['tempo_medio_formatado']])
    writer.writerow(['Índice de Severidade Ponderado', f"{metricas['score_severidade']} ({metricas['severidade_classificacao']})"])
    writer.writerow([])

    # Cabeçalho da tabela de dados
    writer.writerow([
        'ID',
        'Título',
        'Solicitante',
        'Responsável',
        'Status',
        'Prioridade',
        'Situação SLA',
        'Data Criação',
        'Data Limite SLA',
        'Data Conclusão'
    ])

    for d in demandas:
        st_sla = d.get('status_sla', '')
        if st_sla == 'sem_responsavel_fora_sla':
            sit_sla = 'Fora do SLA (Sem Responsável)'
        elif st_sla == 'atrasada':
            sit_sla = 'Atrasada (Fora do SLA)'
        elif st_sla == 'concluida_atrasada':
            sit_sla = 'Concluída com Atraso'
        elif st_sla == 'concluida_no_prazo':
            sit_sla = 'Concluída no Prazo'
        elif st_sla == 'cancelada':
            sit_sla = 'Cancelada'
        else:
            sit_sla = 'No Prazo'

        writer.writerow([
            d.get('id', ''),
            d.get('titulo', ''),
            d.get('solicitante_nome', '') or 'Não informado',
            d.get('responsavel_nome', '') or 'Sem responsável',
            d.get('status', ''),
            d.get('prioridade', ''),
            sit_sla,
            d.get('data_criacao', '') or '',
            d.get('data_limite', '') or '',
            d.get('data_conclusao', '') or '-'
        ])

    # Rodapé com filtros aplicados e timestamp
    writer.writerow([])
    writer.writerow(['------------------------------------------------------------'])
    writer.writerow([f"Filtros aplicados: {metricas['texto_filtros']}"])
    writer.writerow([f"Relatório gerado em: {metricas['timestamp_geracao']}"])

    csv_data = sio.getvalue().encode('utf-8-sig')
    return Response(
        csv_data,
        mimetype='text/csv; charset=utf-8-sig',
        headers={'Content-Disposition': 'attachment; filename=relatorio_demandas_sgdi.csv'}
    )


def gerar_excel_response(metricas):
    """Gera resposta HTTP com planilha Excel (.xlsx) contendo abas Resumo e Detalhes formatadas."""
    demandas = metricas['demandas_lista']
    wb = openpyxl.Workbook()

    # Estilos de formatação
    fonte_titulo = Font(name='Calibri', size=15, bold=True, color='1E293B')
    fonte_subtitulo = Font(name='Calibri', size=11, italic=True, color='64748B')
    fonte_secao = Font(name='Calibri', size=12, bold=True, color='4F46E5')
    fonte_header = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    fonte_dados = Font(name='Calibri', size=10)
    fonte_rodape = Font(name='Calibri', size=9, italic=True, color='64748B')

    fill_header = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    fill_header_u = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    borda_fina = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. ABA RESUMO ANALÍTICO
    ws_resumo = wb.active
    ws_resumo.title = 'Resumo Analítico'

    ws_resumo['A1'] = f'{NOME_EMPRESA} - {NOME_ORGANIZACAO}'
    ws_resumo['A1'].font = fonte_titulo
    ws_resumo['A2'] = 'Painel Executivo de Analytics e SLA'
    ws_resumo['A2'].font = fonte_subtitulo

    ws_resumo['A4'] = 'MÉTRICAS GERAIS CONSOLIDADAS'
    ws_resumo['A4'].font = fonte_secao

    kpi_rows = [
        ('Total de Demandas', metricas['total_demandas']),
        ('Demandas Abertas', metricas['total_abertas']),
        ('Em Andamento', metricas['total_em_andamento']),
        ('Concluídas', metricas['total_concluidas']),
        ('Canceladas', metricas['total_canceladas']),
        ('Fora do SLA / Atrasadas', metricas['total_fora_sla']),
        ('Taxa de Cumprimento de SLA', f"{metricas['taxa_sla_geral']}%"),
        ('Tempo Médio de Resolução (Concluídas)', metricas['tempo_medio_formatado']),
        ('Índice de Severidade Ponderado', f"{metricas['score_severidade']} / 4.0 ({metricas['severidade_classificacao']})"),
    ]

    for idx, (kpi, val) in enumerate(kpi_rows, start=5):
        ws_resumo[f'A{idx}'] = kpi
        ws_resumo[f'B{idx}'] = val
        ws_resumo[f'A{idx}'].font = Font(name='Calibri', bold=True)
        ws_resumo[f'B{idx}'].font = fonte_dados
        ws_resumo[f'A{idx}'].border = borda_fina
        ws_resumo[f'B{idx}'].border = borda_fina

    # Tabela de Responsáveis no Resumo
    linha_u = 16
    ws_resumo[f'A{linha_u}'] = 'DESEMPENHO POR RESPONSÁVEL'
    ws_resumo[f'A{linha_u}'].font = fonte_secao
    linha_u += 1

    headers_resp = ['Responsável', 'Total Atribuído', 'Abertas', 'Em Andamento', 'Concluídas', 'Fora SLA', 'Taxa Conclusão', 'Tempo Médio', 'Carga Ponderada']
    for col_idx, h in enumerate(headers_resp, start=1):
        cell = ws_resumo.cell(row=linha_u, column=col_idx, value=h)
        cell.font = fonte_header
        cell.fill = fill_header_u
        cell.alignment = Alignment(horizontal='center')

    for u in metricas['usuarios_stats']:
        linha_u += 1
        vals = [
            u['nome'], u['total'], u['abertas'], u['em_andamento'], u['concluidas'],
            u['fora_sla'], f"{u['taxa_conclusao']}%", u['tempo_medio_formatado'], f"{u['pontos_criticidade']} pts"
        ]
        for col_idx, val in enumerate(vals, start=1):
            cell = ws_resumo.cell(row=linha_u, column=col_idx, value=val)
            cell.font = fonte_dados
            cell.border = borda_fina
            if col_idx > 1:
                cell.alignment = Alignment(horizontal='center')

    # Rodapé no resumo
    linha_u += 3
    ws_resumo[f'A{linha_u}'] = f"Filtros aplicados: {metricas['texto_filtros']}"
    ws_resumo[f'A{linha_u}'].font = fonte_rodape
    linha_u += 1
    ws_resumo[f'A{linha_u}'] = f"Relatório gerado em: {metricas['timestamp_geracao']}"
    ws_resumo[f'A{linha_u}'].font = fonte_rodape

    # 2. ABA DEMANDAS DETALHADAS
    ws_demandas = wb.create_sheet(title='Lista de Demandas')

    ws_demandas['A1'] = f'{NOME_EMPRESA} - {NOME_ORGANIZACAO}'
    ws_demandas['A1'].font = fonte_titulo
    ws_demandas['A2'] = 'Relatório Detalhado de Demandas e SLA'
    ws_demandas['A2'].font = fonte_subtitulo

    headers_dem = ['ID', 'Título', 'Solicitante', 'Responsável', 'Status', 'Prioridade', 'Situação SLA', 'Criação', 'Prazo Limite SLA', 'Conclusão']
    linha_d = 4
    for col_idx, h in enumerate(headers_dem, start=1):
        cell = ws_demandas.cell(row=linha_d, column=col_idx, value=h)
        cell.font = fonte_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center')

    for d in demandas:
        linha_d += 1
        st_sla = d.get('status_sla', '')
        if st_sla == 'sem_responsavel_fora_sla': sit_sla = 'Fora SLA (Sem Resp.)'
        elif st_sla == 'atrasada': sit_sla = 'Atrasada'
        elif st_sla == 'concluida_atrasada': sit_sla = 'Concluída (Atraso)'
        elif st_sla == 'concluida_no_prazo': sit_sla = 'Concluída no Prazo'
        elif st_sla == 'cancelada': sit_sla = 'Cancelada'
        else: sit_sla = 'No Prazo'

        row_vals = [
            d.get('id', ''),
            d.get('titulo', ''),
            d.get('solicitante_nome', '') or 'Não informado',
            d.get('responsavel_nome', '') or 'Sem responsável',
            d.get('status', ''),
            d.get('prioridade', ''),
            sit_sla,
            d.get('data_criacao', '') or '',
            d.get('data_limite', '') or '',
            d.get('data_conclusao', '') or '-'
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_demandas.cell(row=linha_d, column=col_idx, value=val)
            cell.font = fonte_dados
            cell.border = borda_fina
            if col_idx in (1, 5, 6, 7, 8, 9, 10):
                cell.alignment = Alignment(horizontal='center')

    # Rodapé nas demandas
    linha_d += 2
    ws_demandas[f'A{linha_d}'] = f"Filtros aplicados: {metricas['texto_filtros']}"
    ws_demandas[f'A{linha_d}'].font = fonte_rodape
    linha_d += 1
    ws_demandas[f'A{linha_d}'] = f"Relatório gerado em: {metricas['timestamp_geracao']}"
    ws_demandas[f'A{linha_d}'].font = fonte_rodape

    # Auto-ajuste de colunas
    for ws in (ws_resumo, ws_demandas):
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 3, 40))

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    return send_file(
        bio,
        as_attachment=True,
        download_name='relatorio_demandas_sgdi.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


def gerar_pdf_response(metricas):
    """Gera resposta HTTP com relatório executivo formal em PDF via markdown_pdf."""
    demandas = metricas['demandas_lista']
    md = f"""# {NOME_EMPRESA}
## {NOME_ORGANIZACAO}
### Relatório Executivo de Analytics e Governança de SLA

---

#### 📌 Resumo Geral de Métricas
- **Total de Demandas:** {metricas['total_demandas']}
- **Abertas:** {metricas['total_abertas']} ({metricas['pct_abertas']}%) | **Em Andamento:** {metricas['total_em_andamento']} ({metricas['pct_em_andamento']}%)
- **Concluídas:** {metricas['total_concluidas']} ({metricas['pct_concluidas']}%) | **Canceladas:** {metricas['total_canceladas']} ({metricas['pct_canceladas']}%)
- **Fora do SLA / Atrasadas:** {metricas['total_fora_sla']} (SLA Global: {metricas['taxa_sla_geral']}%)
- **Tempo Médio de Resolução:** {metricas['tempo_medio_formatado']} *(Apenas Concluídas; Canceladas desconsideradas)*
- **Índice de Severidade do Backlog:** {metricas['score_severidade']} / 4.0 ({metricas['severidade_classificacao']})

---

#### 👥 Desempenho por Responsável
| Responsável | Total | Abertas | Em Andamento | Concluídas | Fora SLA | Taxa Conclusão | Tempo Médio |
|---|---|---|---|---|---|---|---|
"""
    for u in metricas['usuarios_stats']:
        md += f"| {u['nome']} | {u['total']} | {u['abertas']} | {u['em_andamento']} | {u['concluidas']} | {u['fora_sla']} | {u['taxa_conclusao']}% | {u['tempo_medio_formatado']} |\n"

    md += """
---

#### 📋 Lista de Demandas
| ID | Título | Solicitante | Responsável | Status | Prioridade | Situação SLA | Criação |
|---|---|---|---|---|---|---|---|
"""
    for d in demandas[:60]:  # Limite de apresentação para PDF limpo
        st_sla = d.get('status_sla', '')
        if st_sla == 'sem_responsavel_fora_sla': sit_sla = 'Fora SLA (Sem Resp)'
        elif st_sla == 'atrasada': sit_sla = 'Atrasada'
        elif st_sla == 'concluida_atrasada': sit_sla = 'Concluída (Atraso)'
        elif st_sla == 'concluida_no_prazo': sit_sla = 'Concluída no Prazo'
        elif st_sla == 'cancelada': sit_sla = 'Cancelada'
        else: sit_sla = 'No Prazo'

        tit = (d.get('titulo', '')[:28] + '...') if len(d.get('titulo', '')) > 28 else d.get('titulo', '')
        solic = d.get('solicitante_nome', '') or '-'
        resp = d.get('responsavel_nome', '') or 'Sem resp.'
        dt_c = (d.get('data_criacao', '')[:10]) if d.get('data_criacao') else '-'
        md += f"| #{d.get('id')} | {tit} | {solic} | {resp} | {d.get('status')} | {d.get('prioridade')} | {sit_sla} | {dt_c} |\n"

    md += f"""
---

**Filtros aplicados:** {metricas['texto_filtros']}  
**Relatório gerado em:** {metricas['timestamp_geracao']}  
*{NOME_EMPRESA} - Confidencial*
"""

    pdf = MarkdownPdf()
    pdf.add_section(Section(md))
    bio = io.BytesIO()
    pdf.save_bytes(bio)
    bio.seek(0)

    return send_file(
        bio,
        as_attachment=True,
        download_name='relatorio_demandas_sgdi.pdf',
        mimetype='application/pdf'
    )
